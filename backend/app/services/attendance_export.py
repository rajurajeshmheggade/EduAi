"""Attendance exports (Excel / CSV) for teachers and admins.

Input is a list of class-subject registers built by `attendance.monthly_register` over any date range.
Three report shapes:
  register   - day-wise P / A / L / LV grid per class-subject (one sheet each in Excel)
  summary    - one row per student per class-subject with totals and %
  defaulters - the summary filtered to students below the institution threshold
Excel always starts with an "Overview" sheet (period, threshold and a per class-subject roll-up).
"""
from __future__ import annotations

import csv
import io
import re
from datetime import date
from typing import Literal

Report = Literal["register", "summary", "defaulters"]
ABBREV = {"present": "P", "absent": "A", "late": "L", "excused": "LV"}
TOTAL_HEADERS = ["Present", "Absent", "Late", "Leave", "Countable", "Attendance %", "Status"]


def _day_label(d: str) -> str:
    return date.fromisoformat(d).strftime("%d-%b")


def _cell(statuses: list[str]) -> str:
    return "/".join(ABBREV[s] for s in statuses)


def _totals(s: dict, threshold: float) -> list:
    countable = s["total"] - s["excused"]
    pct = s["percentage"]
    status = "No data" if pct is None else ("Below threshold" if pct < threshold else "OK")
    return [s["present"], s["absent"], s["late"], s["excused"], countable, "" if pct is None else pct, status]


def _student_rows(regs: list[dict], defaulters_only: bool) -> list[list]:
    rows = []
    for r in regs:
        for s in r["reg"]["students"]:
            if defaulters_only and not s["is_defaulter"]:
                continue
            rows.append([r["class_name"], r["subject_name"], s["roll_number"] or "", s["student_user_code"], s["full_name"],
                         *_totals(s, r["reg"]["threshold"])])
    return rows


STUDENT_HEADERS = ["Class", "Subject", "Roll", "Admission No", "Name", *TOTAL_HEADERS]


def _overview_rows(regs: list[dict]) -> list[list]:
    out = []
    for r in regs:
        g = r["reg"]
        st = g["students"]
        p, a, l, lv = (sum(s[k] for s in st) for k in ("present", "absent", "late", "excused"))
        countable = p + a + l
        out.append([r["class_name"], r["subject_name"], g["sessions_held"], len(st), p, a, l, lv,
                    round(100 * (p + l) / countable, 1) if countable else "", len(g["defaulters"])])
    return out


OVERVIEW_HEADERS = ["Class", "Subject", "Periods held", "Students", "Present", "Absent", "Late", "Leave", "Attendance %", "Below threshold"]


# ---------------------------------------------------------------- CSV
def to_csv(regs: list[dict], report: Report) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf)
    if report == "register":
        days = sorted({d for r in regs for d in r["reg"]["marked_days"]})
        w.writerow(["Class", "Subject", "Roll", "Admission No", "Name", *[_day_label(d) for d in days], *TOTAL_HEADERS])
        for r in regs:
            for s in r["reg"]["students"]:
                w.writerow([r["class_name"], r["subject_name"], s["roll_number"] or "", s["student_user_code"], s["full_name"],
                            *[_cell(s["days"].get(d, [])) for d in days], *_totals(s, r["reg"]["threshold"])])
    else:
        w.writerow(STUDENT_HEADERS)
        w.writerows(_student_rows(regs, report == "defaulters"))
    # BOM so Excel opens UTF-8 names correctly.
    return ("﻿" + buf.getvalue()).encode("utf-8")


# ---------------------------------------------------------------- Excel
def _sheet_title(name: str, used: set[str]) -> str:
    base = re.sub(r"[\[\]:*?/\\]", "-", name)[:31] or "Sheet"
    title, i = base, 2
    while title.lower() in used:
        suffix = f" ({i})"
        title, i = base[: 31 - len(suffix)] + suffix, i + 1
    used.add(title.lower())
    return title


def to_xlsx(regs: list[dict], report: Report, meta: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    brand = PatternFill("solid", fgColor="4F46E5")
    head_font = Font(bold=True, color="FFFFFF")
    fills = {"A": PatternFill("solid", fgColor="FDE2E2"), "L": PatternFill("solid", fgColor="FEF3C7"),
             "LV": PatternFill("solid", fgColor="E0E7FF"), "P": PatternFill("solid", fgColor="DCFCE7")}
    low = PatternFill("solid", fgColor="FEE2E2")
    thin = Side(style="thin", color="E5E7EB")
    grid = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center")
    used: set[str] = set()

    def title_block(ws, title: str, extra: list[str]) -> int:
        ws.append([title]); ws["A1"].font = Font(bold=True, size=14, color="1E1B4B")
        for line in [f"{meta['org_name']} · {meta['period']} · Threshold {meta['threshold']:g}%", *extra,
                     f"Generated {meta['generated_at']} by {meta['generated_by']}"]:
            ws.append([line]); ws.cell(ws.max_row, 1).font = Font(color="6B7280")
        ws.append([])
        return ws.max_row + 1

    def header(ws, row: int, headers: list[str]) -> None:
        for c, h in enumerate(headers, 1):
            cell = ws.cell(row, c, h)
            cell.fill, cell.font, cell.alignment, cell.border = brand, head_font, center, grid

    def body(ws, start: int, rows: list[list], pct_col: int, status_col: int | None = None) -> None:
        for i, values in enumerate(rows):
            r = start + i
            for c, v in enumerate(values, 1):
                cell = ws.cell(r, c, v)
                cell.border = grid
            pct = values[pct_col - 1]
            if isinstance(pct, (int, float)) and pct < meta["threshold"]:
                ws.cell(r, pct_col).fill = low
                ws.cell(r, pct_col).font = Font(bold=True, color="B91C1C")
                if status_col:
                    ws.cell(r, status_col).font = Font(bold=True, color="B91C1C")

    def widths(ws, first_cols: dict[int, int], default: int = 9) -> None:
        for c in range(1, ws.max_column + 1):
            ws.column_dimensions[get_column_letter(c)].width = first_cols.get(c, default)

    wb = Workbook()
    ov = wb.active
    ov.title = _sheet_title("Overview", used)
    start = title_block(ov, "Attendance overview", [f"Report: {report.title()}"])
    header(ov, start, OVERVIEW_HEADERS)
    body(ov, start + 1, _overview_rows(regs), pct_col=9)
    widths(ov, {1: 22, 2: 22, 3: 13, 4: 10, 9: 13, 10: 15})
    ov.freeze_panes = ov.cell(start + 1, 1)
    legend = ov.max_row + 2
    ov.cell(legend, 1, "Legend").font = Font(bold=True)
    for i, (k, label) in enumerate([("P", "Present"), ("A", "Absent"), ("L", "Late"), ("LV", "Leave (excused)")], 1):
        c = ov.cell(legend + i, 1, k); c.fill, c.alignment = fills[k], center
        ov.cell(legend + i, 2, label)
    ov.cell(legend + 6, 1, "Attendance % = (present + late) / (periods marked − leave) × 100").font = Font(italic=True, color="6B7280")

    if report == "register":
        for r in regs:
            g = r["reg"]
            ws = wb.create_sheet(_sheet_title(f"{r['class_code']} {r['subject_code']}", used))
            avg = "—" if g["class_average"] is None else f"{g['class_average']}%"
            start = title_block(ws, f"{r['class_name']} · {r['subject_name']}", [f"Periods held {g['sessions_held']} · Class average {avg}"])
            days = g["marked_days"]
            heads = ["Roll", "Admission No", "Name", *[_day_label(d) for d in days], *TOTAL_HEADERS]
            header(ws, start, heads)
            rows = [[s["roll_number"] or "", s["student_user_code"], s["full_name"], *[_cell(s["days"].get(d, [])) for d in days],
                     *_totals(s, g["threshold"])] for s in g["students"]]
            pct_col = 3 + len(days) + 6
            body(ws, start + 1, rows, pct_col=pct_col, status_col=pct_col + 1)
            for i in range(len(rows)):
                for c in range(4, 4 + len(days)):
                    cell = ws.cell(start + 1 + i, c)
                    cell.alignment = center
                    first = str(cell.value or "").split("/")[0]
                    if first in fills and first != "P":
                        cell.fill = fills[first]
            # Day totals row.
            tr = start + 1 + len(rows)
            ws.cell(tr, 3, "Present / marked").font = Font(bold=True)
            for j, d in enumerate(days):
                t = g["day_totals"][d]
                ws.cell(tr, 4 + j, f"{t['present'] + t['late']}/{t['total'] - t['excused']}").alignment = center
            widths(ws, {1: 7, 2: 14, 3: 26, **{c: 7 for c in range(4, 4 + len(days))}, pct_col: 13, pct_col + 1: 15})
            ws.freeze_panes = ws.cell(start + 1, 4)
    else:
        ws = wb.create_sheet(_sheet_title("Defaulters" if report == "defaulters" else "Students", used))
        start = title_block(ws, "Students below threshold" if report == "defaulters" else "Student-wise attendance", [])
        header(ws, start, STUDENT_HEADERS)
        rows = _student_rows(regs, report == "defaulters")
        body(ws, start + 1, rows, pct_col=11, status_col=12)
        if not rows:
            ws.cell(start + 1, 1, "No students below the threshold in this period." if report == "defaulters" else "No students.")
        widths(ws, {1: 18, 2: 20, 3: 7, 4: 14, 5: 26, 10: 11, 11: 13, 12: 15})
        ws.freeze_panes = ws.cell(start + 1, 1)
        ws.auto_filter.ref = f"A{start}:{get_column_letter(len(STUDENT_HEADERS))}{start + max(len(rows), 1)}"

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
