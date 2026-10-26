from datetime import date, timedelta

import pytest

from app.services import attendance as att

A = "/ICMELive/api/v1/admin"
C = "/ICMELive/api/v1/classroom"
TODAY = date.today().isoformat()
MONTH = date.today().strftime("%Y-%m")


def _login(client, ident, password):
    r = client.post("/api/v1/auth/login", json={"email": ident, "password": password})
    assert r.status_code == 200, r.text
    return {"Authorization": f"Bearer {r.json()['access_token']}"}


@pytest.fixture(scope="module")
def superadmin(client):
    return _login(client, "superadmin", "superadmin")


@pytest.fixture(scope="module")
def tenant(client, superadmin):
    """A fresh institution: admin, class teacher, subject teacher, one class with two subjects."""
    r = client.post(f"{A}/organizations", headers=superadmin, json={
        "org_code": "GREENFIELD", "org_name": "Greenfield College", "institution_level": "college",
        "admin": {"full_name": "Gita Admin", "email": "gita@greenfield.edu", "user_code": "GF-ADM1"},
    })
    assert r.status_code == 201, r.text
    admin = _login(client, "gita@greenfield.edu", r.json()["admin"]["temporary_password"])

    def teacher(code, name, email):
        t = client.post(f"{A}/teachers", headers=admin, json={"user_code": code, "full_name": name, "email": email,
                                                              "password": "Teach3r!pass"})
        assert t.status_code == 201, t.text
        assert t.json()["temporary_password"] is None  # admin chose the password
        return code

    teacher("GF-T1", "Ravi Class Teacher", "ravi@greenfield.edu")
    teacher("GF-T2", "Sana Subject Teacher", "sana@greenfield.edu")
    c = client.post(f"{A}/classes", headers=admin, json={"class_code": "BSC-1", "class_name": "B.Sc Year 1",
                                                         "class_teacher_code": "GF-T1"})
    assert c.status_code == 201, c.text
    for subj, name, t in [("CHEM-1", "Chemistry", "GF-T2"), ("BIO-1", "Biology", "GF-T1")]:
        r = client.put(f"{A}/classes/BSC-1/subjects/{subj}", headers=admin, json={"subject_name": name, "teacher_code": t})
        assert r.status_code == 200, r.text
    return {
        "admin": admin,
        "ravi": _login(client, "ravi@greenfield.edu", "Teach3r!pass"),
        "sana": _login(client, "GF-T2", "Teach3r!pass"),  # login by employee code
    }


@pytest.fixture(scope="module")
def students(client, tenant):
    out = {}
    for code, name, roll in [("GF-S1", "Asha", "1"), ("GF-S2", "Bilal", "2"), ("GF-S3", "Chen", "3")]:
        r = client.post(f"{C}/classes/BSC-1/students", headers=tenant["ravi"],
                        json={"full_name": name, "user_code": code, "roll_number": roll})
        assert r.status_code == 201, r.text
        out[code] = r.json()
    return out


# ---------------------------------------------------------------- unit: attendance maths
def test_tally_counts_late_as_attended_and_ignores_excused():
    t = att.tally(["present", "late", "absent", "excused"])
    assert (t["present"], t["late"], t["absent"], t["excused"], t["total"]) == (1, 1, 1, 1, 4)
    assert t["percentage"] == pytest.approx(66.7)
    assert att.tally(["excused"])["percentage"] is None
    assert att.tally([])["percentage"] is None


def test_month_bounds():
    assert att.month_bounds("2028-02") == (date(2028, 2, 1), date(2028, 2, 29))
    with pytest.raises(ValueError):
        att.month_bounds("2026-13")


# ---------------------------------------------------------------- hierarchy & RBAC
def test_only_superadmin_onboards_institutions(client, admin, teacher):
    body = {"org_code": "X1", "org_name": "Nope", "admin": {"full_name": "N A", "email": "n@x.edu", "user_code": "NX1"}}
    assert client.post(f"{A}/organizations", headers=admin, json=body).status_code == 403
    assert client.post(f"{A}/organizations", headers=teacher, json=body).status_code == 403


def test_superadmin_lists_tenants_with_counts(client, superadmin, tenant):
    orgs = {o["org_code"]: o for o in client.get(f"{A}/organizations", headers=superadmin).json()["organizations"]}
    assert orgs["GREENFIELD"]["teachers"] == 2 and orgs["GREENFIELD"]["classes"] == 1


def test_teachers_and_students_cannot_use_admin_api(client, tenant, teacher, student):
    for h in (teacher, student, tenant["ravi"]):
        assert client.post(f"{A}/teachers", headers=h, json={"user_code": "Z", "full_name": "Zed", "email": "z@z.io"}).status_code == 403
        assert client.post(f"{A}/classes", headers=h, json={"class_code": "Z", "class_name": "Zed"}).status_code == 403


def test_duplicate_teacher_is_409(client, tenant):
    r = client.post(f"{A}/teachers", headers=tenant["admin"],
                    json={"user_code": "GF-T1", "full_name": "Dup", "email": "dup@greenfield.edu"})
    assert r.status_code == 409


def test_class_teacher_must_be_a_teacher_of_this_institution(client, tenant):
    r = client.post(f"{A}/classes", headers=tenant["admin"],
                    json={"class_code": "BSC-2", "class_name": "B.Sc Year 2", "class_teacher_code": "EMP1001"})
    assert r.status_code == 422  # EMP1001 belongs to DEMO-SCHOOL


def test_subject_assignment_flows_into_teacher_subjects(client, tenant):
    me = client.get("/api/v1/auth/me", headers=tenant["sana"]).json()
    assert me["allowed_subjects"] == ["CHEM-1"]
    cls = client.get(f"{A}/classes", headers=tenant["admin"]).json()["classes"][0]
    assert cls["class_teacher_name"] == "Ravi Class Teacher"
    assert {s["subject_code"]: s["teacher_code"] for s in cls["subjects"]} == {"BIO-1": "GF-T1", "CHEM-1": "GF-T2"}


def test_tenant_isolation(client, tenant, admin, teacher):
    assert all(c["class_code"] != "BSC-1" for c in client.get(f"{A}/classes", headers=admin).json()["classes"])
    assert client.get(f"{C}/classes/BSC-1/students", headers=teacher).status_code == 404


# ---------------------------------------------------------------- students: class teacher only
def test_class_teacher_creates_students_who_get_class_subjects(client, students):
    s = students["GF-S1"]
    assert s["temporary_password"] and s["email"].endswith(".noemail.local")
    h = _login(client, "GF-S1", s["temporary_password"])  # login with admission number
    me = client.get("/api/v1/auth/me", headers=h).json()
    assert me["role"] == "student" and me["class_code"] == "BSC-1"
    assert sorted(me["allowed_subjects"]) == ["BIO-1", "CHEM-1"]


def test_subject_teacher_and_admin_cannot_create_students(client, tenant, students):
    body = {"full_name": "Intruder"}
    assert client.post(f"{C}/classes/BSC-1/students", headers=tenant["sana"], json=body).status_code == 403
    assert client.post(f"{C}/classes/BSC-1/students", headers=tenant["admin"], json=body).status_code == 403
    # but the subject teacher can see the roster of a class she teaches
    r = client.get(f"{C}/classes/BSC-1/students", headers=tenant["sana"]).json()
    assert r["is_class_teacher"] is False and len(r["students"]) == 3


def test_csv_import_reports_bad_rows(client, tenant, students):
    csv = "full_name,user_code,roll_number,email\nDev Rao,GF-S4,4,dev@greenfield.edu\n,GF-S5,5,\nDup,GF-S1,6,\n"
    r = client.post(f"{C}/classes/BSC-1/students/import", headers=tenant["ravi"],
                    files={"file": ("s.csv", csv.encode(), "text/csv")})
    assert r.status_code == 200, r.text
    body = r.json()
    assert [c["user_code"] for c in body["created"]] == ["GF-S4"]
    assert [e["row"] for e in body["errors"]] == [3, 4]
    # deactivate the imported student so later rosters stay at three
    p = client.patch(f"{C}/classes/BSC-1/students/GF-S4", headers=tenant["ravi"], json={"is_active": False})
    assert p.status_code == 200 and p.json()["is_active"] is False


def test_change_password(client, students):
    tmp = students["GF-S2"]["temporary_password"]
    h = _login(client, "GF-S2", tmp)
    assert client.post("/api/v1/auth/change-password", headers=h,
                       json={"current_password": "wrong", "new_password": "NewPass123"}).status_code == 400
    assert client.post("/api/v1/auth/change-password", headers=h,
                       json={"current_password": tmp, "new_password": "NewPass123"}).status_code == 200
    _login(client, "GF-S2", "NewPass123")


# ---------------------------------------------------------------- attendance marking
def _mark(client, h, subject, statuses, slot="Period-1", day=TODAY):
    recs = [{"student_user_code": s, "status": st} for s, st in statuses.items()]
    return client.post(f"{C}/attendance", headers=h, json={"class_code": "BSC-1", "subject_code": subject,
                                                           "attendance_date": day, "session_slot": slot,
                                                           "records": recs})


def test_subject_teacher_marks_only_own_subject(client, tenant, students):
    r = _mark(client, tenant["sana"], "CHEM-1", {"GF-S1": "present", "GF-S2": "absent", "GF-S3": "late"})
    assert r.status_code == 200 and r.json()["saved"] == 3
    assert _mark(client, tenant["sana"], "BIO-1", {"GF-S1": "present"}).status_code == 403


def test_class_teacher_can_mark_any_subject_of_class(client, tenant, students):
    assert _mark(client, tenant["ravi"], "CHEM-1", {"GF-S1": "present", "GF-S2": "present", "GF-S3": "present"},
                 slot="Period-2").status_code == 200
    assert _mark(client, tenant["ravi"], "BIO-1", {"GF-S1": "absent", "GF-S2": "excused", "GF-S3": "present"}).status_code == 200


def test_attendance_validation(client, tenant, students, teacher):
    tomorrow = (date.today() + timedelta(days=1)).isoformat()
    assert _mark(client, tenant["ravi"], "BIO-1", {"GF-S1": "present"}, day=tomorrow).status_code == 422
    assert _mark(client, tenant["ravi"], "BIO-1", {"REG2026001": "present"}).status_code == 422  # not in class
    assert _mark(client, tenant["ravi"], "BIO-1", {"GF-S4": "present"}).status_code == 422  # deactivated
    assert _mark(client, tenant["ravi"], "BIO-1", {"GF-S1": "holiday"}).status_code == 422
    assert _mark(client, teacher, "BIO-1", {"GF-S1": "present"}).status_code == 404  # other tenant


def test_remarking_updates_instead_of_duplicating(client, tenant, students):
    assert _mark(client, tenant["sana"], "CHEM-1", {"GF-S2": "late"}).status_code == 200
    day = client.get(f"{C}/attendance", headers=tenant["sana"],
                     params={"class_code": "BSC-1", "subject_code": "CHEM-1", "attendance_date": TODAY}).json()
    assert day["is_marked"] and {s["student_user_code"]: s["status"] for s in day["students"]} == {
        "GF-S1": "present", "GF-S2": "late", "GF-S3": "late"}


def test_monthly_register_and_csv(client, tenant, students):
    q = {"class_code": "BSC-1", "subject_code": "CHEM-1", "month": MONTH}
    reg = client.get(f"{C}/attendance/register", headers=tenant["sana"], params=q).json()
    rows = {s["student_user_code"]: s for s in reg["students"]}
    assert reg["sessions_held"] == 2 and reg["marked_days"] == [TODAY]
    assert rows["GF-S1"]["percentage"] == 100.0 and rows["GF-S2"]["percentage"] == 100.0  # late counts as attended
    assert rows["GF-S2"]["days"][TODAY] == ["late", "present"] or rows["GF-S2"]["days"][TODAY] == ["present", "late"]
    assert reg["threshold"] == 75.0
    # admin can view the register institution-wide; a different subject teacher cannot
    assert client.get(f"{C}/attendance/register", headers=tenant["admin"], params=q).status_code == 200
    bio = {**q, "subject_code": "BIO-1"}
    assert client.get(f"{C}/attendance/register", headers=tenant["sana"], params=bio).status_code == 403
    reg_bio = client.get(f"{C}/attendance/register", headers=tenant["ravi"], params=bio).json()
    assert reg_bio["defaulters"] == ["GF-S1"]  # absent in the only session; S2 excused -> no percentage

    csv = client.get(f"{C}/attendance/register.csv", headers=tenant["ravi"], params=q)
    assert csv.status_code == 200 and csv.headers["content-type"].startswith("text/csv")
    lines = csv.text.strip().splitlines()
    assert lines[0].startswith("Roll,Admission No,Name") and len(lines) == 4


def test_attendance_export_xlsx_and_csv(client, tenant, students):
    import io

    from openpyxl import load_workbook

    rng = {"start": (date.today() - timedelta(days=40)).isoformat(), "end": TODAY}
    # class teacher: whole class (both subjects), register workbook
    r = client.get(f"{C}/attendance/export", headers=tenant["ravi"], params={**rng, "class_code": "BSC-1"})
    assert r.status_code == 200, r.text
    assert r.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert ".xlsx" in r.headers["content-disposition"]
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Overview", "BSC-1 BIO-1", "BSC-1 CHEM-1"]
    chem = [[c for c in row] for row in wb["BSC-1 CHEM-1"].iter_rows(values_only=True)]
    names = [row[2] for row in chem if row and row[1] in ("GF-S1", "GF-S2", "GF-S3")]
    assert names == ["Asha", "Bilal", "Chen"]

    # subject teacher: only their subject; whole class export narrows to it, other subject is 403
    r = client.get(f"{C}/attendance/export", headers=tenant["sana"], params={**rng, "class_code": "BSC-1", "fmt": "csv"})
    assert r.status_code == 200 and "Biology" not in r.text and "Chemistry" in r.text
    assert r.content.startswith(b"\xef\xbb\xbf")  # BOM for Excel
    assert client.get(f"{C}/attendance/export", headers=tenant["sana"],
                      params={**rng, "class_code": "BSC-1", "subject_code": "BIO-1"}).status_code == 403
    # a teacher must pick a class
    assert client.get(f"{C}/attendance/export", headers=tenant["sana"], params=rng).status_code == 422

    # admin: whole institution, defaulters only (GF-S1 absent in Biology)
    r = client.get(f"{C}/attendance/export", headers=tenant["admin"], params={**rng, "report": "defaulters", "fmt": "csv"})
    assert r.status_code == 200, r.text
    lines = r.content.decode("utf-8-sig").strip().splitlines()
    assert lines[0].startswith("Class,Subject,Roll,Admission No,Name")
    assert len(lines) == 2 and "GF-S1" in lines[1] and "Biology" in lines[1] and "Below threshold" in lines[1]
    r = client.get(f"{C}/attendance/export", headers=tenant["admin"], params={**rng, "report": "summary"})
    assert r.status_code == 200 and load_workbook(io.BytesIO(r.content)).sheetnames == ["Overview", "Students"]

    # validation
    assert client.get(f"{C}/attendance/export", headers=tenant["admin"],
                      params={"start": TODAY, "end": "2000-01-01"}).status_code == 422
    assert client.get(f"{C}/attendance/export", headers=tenant["admin"],
                      params={"start": "2020-01-01", "end": TODAY}).status_code == 422


def test_register_any_range_lists_every_day():
    reg = att.monthly_register([], [], date(2026, 1, 30), date(2026, 2, 2), 75)
    assert reg["days"] == ["2026-01-30", "2026-01-31", "2026-02-01", "2026-02-02"]


def test_student_sees_own_subject_wise_monthly_attendance(client, tenant, students):
    h = _login(client, "GF-S1", students["GF-S1"]["temporary_password"])
    r = client.get(f"{C}/my-attendance", headers=h, params={"month": MONTH})
    assert r.status_code == 200, r.text
    body = r.json()
    subj = {s["subject_code"]: s for s in body["subjects"]}
    assert set(subj) == {"BIO-1", "CHEM-1"}
    assert subj["CHEM-1"]["subject_name"] == "Chemistry"
    assert subj["CHEM-1"]["month"]["present"] == 2 and subj["CHEM-1"]["month"]["percentage"] == 100.0
    assert subj["BIO-1"]["month"]["absent"] == 1 and subj["BIO-1"]["is_below_threshold"] is True
    assert subj["CHEM-1"]["trend"][-1]["month"] == MONTH
    assert all(rec["date"] == TODAY for rec in subj["CHEM-1"]["records"])
    assert body["summary"]["total"] == 3
    # other months are empty, not errors
    old = client.get(f"{C}/my-attendance", headers=h, params={"month": "2020-01"}).json()
    assert all(s["month"]["total"] == 0 for s in old["subjects"])


def test_my_attendance_is_student_only(client, tenant):
    assert client.get(f"{C}/my-attendance", headers=tenant["ravi"]).status_code == 403


def test_admin_overview_and_summary(client, tenant, students):
    ov = client.get(f"{A}/overview", headers=tenant["admin"]).json()
    assert ov["counts"]["teachers"] == 2 and ov["counts"]["students"] == 3 and ov["counts"]["classes"] == 1
    assert ov["attendance_today"]["total"] == 9
    summ = client.get(f"{A}/attendance/summary", headers=tenant["admin"], params={"month": MONTH}).json()
    assert {(r["class_code"], r["subject_code"]) for r in summ["rows"]} == {("BSC-1", "BIO-1"), ("BSC-1", "CHEM-1")}


def test_removing_subject_updates_students_and_teacher(client, tenant, students):
    r = client.delete(f"{A}/classes/BSC-1/subjects/BIO-1", headers=tenant["admin"])
    assert r.status_code == 200
    assert client.get("/api/v1/auth/me", headers=tenant["ravi"]).json()["allowed_subjects"] == []
    st = {s["user_code"]: s for s in client.get(f"{A}/students", headers=tenant["admin"], params={"class_code": "BSC-1"}).json()["students"]}
    assert st["GF-S1"]["allowed_subjects"] == ["CHEM-1"]


def test_suspended_institution_cannot_log_in(client, superadmin, tenant):
    r = client.put(f"{A}/settings", headers=tenant["admin"], json={"is_active": False})
    assert r.status_code == 403  # only the platform owner can suspend
    assert client.put(f"{A}/settings", headers=superadmin, params={"org_code": "GREENFIELD"},
                      json={"is_active": False}).status_code == 200
    assert client.post("/api/v1/auth/login", json={"email": "ravi@greenfield.edu", "password": "Teach3r!pass"}).status_code == 401
    client.put(f"{A}/settings", headers=superadmin, params={"org_code": "GREENFIELD"}, json={"is_active": True})
